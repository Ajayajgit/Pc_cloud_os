# -*- coding: utf-8 -*-
"""Render an assessment draft (markdown) into a branded .docx.

Canonical renderer, consolidated 2026-08-16 from ACFB/working/render_assessment.py
(TOC field, bookmarks, appendix hyperlinks, updateFields) and Teppermans Rapid
Assessment/working/render_docx.py (tier-tag coloring, zebra striping, heading rules,
footer). Brand values follow POWERCLOUD_BRAND_GUIDELINES.md.

Run from the client folder:

    python "../.claude/scripts/render_docx.py" --md working/draft_assessment.md ^
        --out "deliverables/Client_Assessment_2026-01-01.docx" ^
        --client "Client Name" --client-logo source/Logo.png ^
        --cover-line "Commercial Rapid Assessment (Onsite)" ^
        --cover-line "Onsite 2026-01-01 to 2026-01-02"

Markdown subset understood: ## / ### / #### headings (bookmarked, feed the TOC),
tables, - bullets, 1. numbered lists, **bold**, `code`, [text](#anchor) internal
links (anchor slugified the same way as headings, so [OC-1](#oc-1-title) hits the
"#### OC-1 title" entry), [Verified]/[Reported]/[Inferred] tier tags (colored),
*[italic note]* lines, --- rules (skipped).

The TOC is a real field: it populates when export_pdf.py updates fields, or on
open in Word (updateFields is set). Confirm every TOC entry and appendix link
resolves in the exported PDF before delivering.
"""
import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OS_ROOT = Path(__file__).resolve().parents[2]
PC_LOGO = OS_ROOT / "logos" / "Powercloud_Logo-01.png"

SKY_BLUE = RGBColor(0x31, 0xA4, 0xF3)
DEEP_BLUE = RGBColor(0x00, 0x2F, 0x59)
GRASS_GREEN = RGBColor(0x41, 0xB9, 0x46)
MATTE_BLACK = RGBColor(0x1E, 0x1E, 0x1E)
GREY = RGBColor(0x6E, 0x6E, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ZEBRA_FILL = 'EAF4FC'
HEAD, BODY, MONO = 'Montserrat', 'Barlow', 'Consolas'

TIER_RE = re.compile(r'(\[(?:Verified|Reported|Inferred)[^\]]*\])')
INLINE_RE = re.compile(r'\*\*(.+?)\*\*|`([^`]+)`|\[([^\]]+)\]\(#([^)]+)\)')

bookmark_counter = [0]


def tier_color(tok):
    if tok.startswith('[Verified'):
        return GRASS_GREEN
    if tok.startswith('[Reported'):
        return SKY_BLUE
    return GREY


def slugify(anchor_text):
    return 'bm_' + re.sub(r'[^a-z0-9]+', '_', anchor_text.lower()).strip('_')


def add_bookmark(paragraph, name):
    bookmark_counter[0] += 1
    bid = str(bookmark_counter[0])
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), bid)
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink_run(paragraph, text, bookmark_name):
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), bookmark_name)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), BODY)
    rFonts.set(qn('w:hAnsi'), BODY)
    rPr.append(rFonts)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '31A4F3')
    rPr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    rPr.append(OxmlElement('w:b'))
    run.append(rPr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    begin.set(qn('w:dirty'), 'true')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    run._r.append(begin)
    run._r.append(instr)
    sep = paragraph.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    sep._r.append(fld_sep)
    placeholder = paragraph.add_run(
        "Right-click and choose Update Field (or press F9) to populate this table of contents.")
    placeholder.italic = True
    end = paragraph.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    end._r.append(fld_end)


def set_updatefields(document):
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    document.settings.element.append(uf)


def style_run(run, font=BODY, size=10.5, bold=False, italic=False, color=MATTE_BLACK):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), font)


def add_para(doc, text='', font=BODY, size=10.5, bold=False, italic=False,
             color=MATTE_BLACK, space_after=6, align=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        style_run(p.add_run(text), font=font, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc, text, level, bookmark=None):
    """level 1..3 mapped from ## / ### / ####. Sets the outline level so the TOC field picks it up."""
    sizes = {1: 18, 2: 13.5, 3: 11.5}
    colors = {1: DEEP_BLUE, 2: SKY_BLUE, 3: DEEP_BLUE}
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 3)
    pPr = p._p.get_or_add_pPr()
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), str(level - 1))
    pPr.append(outline)
    style_run(p.add_run(text), font=HEAD, size=sizes.get(level, 11),
              bold=True, color=colors.get(level, DEEP_BLUE))
    if level == 1:  # accent rule under top-level headings
        pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '2'), ('w:color', '31A4F3')):
            bottom.set(qn(k), v)
        pbdr.append(bottom)
        pPr.append(pbdr)
    if bookmark:
        add_bookmark(p, bookmark)
    return p


def add_rich_text(paragraph, text, base_font=BODY, base_size=10.5, base_color=MATTE_BLACK,
                  base_bold=False):
    """Inline markdown: **bold**, `code`, [text](#anchor), tier tags."""
    # unescape a literal "\|" (needed inside a markdown table cell to avoid being read
    # as a column separator; harmless no-op here for text that never had one, e.g.
    # already-restored table cells reaching this function a second time)
    text = text.replace('\\|', '|')
    pos = 0
    tokens = []
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            tokens.append(('text', text[pos:m.start()]))
        if m.group(1) is not None:
            tokens.append(('bold', m.group(1)))
        elif m.group(2) is not None:
            tokens.append(('code', m.group(2)))
        else:
            tokens.append(('link', m.group(3), m.group(4)))
        pos = m.end()
    if pos < len(text):
        tokens.append(('text', text[pos:]))

    for tok in tokens:
        if tok[0] == 'link':
            add_internal_hyperlink_run(paragraph, tok[1], slugify(tok[2]))
            continue
        if tok[0] == 'code':
            style_run(paragraph.add_run(tok[1]), font=MONO, size=base_size - 0.5, color=DEEP_BLUE)
            continue
        bold = base_bold or tok[0] == 'bold'
        # a **bold** token may itself wrap a `code` span (e.g. "**`field_id`**: ..."),
        # which the single-pass tokenizer above doesn't see; split it out here so it
        # still renders monospace instead of literal backticks.
        parts = re.split(r'(`[^`]+`)', tok[1]) if tok[0] == 'bold' else [tok[1]]
        for part in parts:
            if not part:
                continue
            if tok[0] == 'bold' and part.startswith('`') and part.endswith('`'):
                style_run(paragraph.add_run(part[1:-1]), font=MONO, size=base_size - 0.5,
                          bold=True, color=DEEP_BLUE)
                continue
            # split tier tags out of the remaining text so they get their color
            for seg in TIER_RE.split(part):
                if not seg:
                    continue
                r = paragraph.add_run(seg)
                if TIER_RE.match(seg):
                    style_run(r, font=base_font, size=base_size, bold=True, color=tier_color(seg))
                else:
                    style_run(r, font=base_font, size=base_size, bold=bold, color=base_color)


def parse_markdown_table(lines):
    rows = []
    for ln in lines:
        ln = ln.strip()
        if not ln.startswith('|'):
            continue
        # protect escaped pipes (literal "|" inside a cell, e.g. "PC \| Script Name")
        # from being read as column separators, then restore them after splitting
        protected = ln.replace('\\|', '\x00')
        cells = [c.strip().replace('\x00', '|') for c in protected.strip('|').split('|')]
        if all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc, rows):
    if not rows:
        return None
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for ridx, row in enumerate(rows):
        for cidx, cell_text in enumerate(row):
            if cidx >= len(table.rows[ridx].cells):
                continue
            cell = table.cell(ridx, cidx)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if ridx == 0:
                add_rich_text(p, cell_text, base_font=HEAD, base_size=9.5, base_color=WHITE,
                              base_bold=True)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:val'), 'clear')
                shading.set(qn('w:fill'), '002F59')
                cell._tc.get_or_add_tcPr().append(shading)
            else:
                if ridx % 2 == 0:  # zebra stripe every other data row
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:val'), 'clear')
                    shading.set(qn('w:fill'), ZEBRA_FILL)
                    cell._tc.get_or_add_tcPr().append(shading)
                add_rich_text(p, cell_text, base_size=9.0)
    # repeat the header row across page breaks
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_code_block(doc, code_lines):
    """A fenced ``` block: literal monospace lines, no inline markdown parsing.
    Used for ASCII diagrams and similar technical figures in TDD-style docs."""
    if not code_lines:
        return
    for line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run(line if line else ' ')
        run.font.name = MONO
        run.font.size = Pt(9)
        run.font.color.rgb = DEEP_BLUE
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
            rFonts.set(qn(attr), MONO)
        t = run._element.find(qn('w:t'))
        if t is not None:
            t.set(qn('xml:space'), 'preserve')
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_cover(doc, args):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    for _ in range(3):
        doc.add_paragraph()
    client_logo = Path(args.client_logo) if args.client_logo else None
    if client_logo and client_logo.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(client_logo), width=Inches(3.2))
        doc.add_paragraph()
        add_para(doc, "in partnership with", size=10, italic=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    if PC_LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(PC_LOGO), width=Inches(1.8))
    doc.add_paragraph()
    add_para(doc, args.client, font=HEAD, size=28, bold=True, color=DEEP_BLUE,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, args.doc_title, font=HEAD, size=18, bold=True, color=SKY_BLUE,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
    for line in args.cover_line or []:
        add_para(doc, line, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(doc, "Prepared by Power Cloud Consulting", size=11,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)
    add_para(doc, "CONFIDENTIAL", font=HEAD, size=12, bold=True, color=GRASS_GREEN,
             align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_footer(doc, args):
    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(fp.add_run(f"Power Cloud Consulting   |   {args.client} {args.doc_title}   |   Confidential"),
              size=8, color=GREY)


def add_toc_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    style_run(p.add_run("Table of Contents"), font=HEAD, size=18, bold=True, color=DEEP_BLUE)
    add_toc_field(doc.add_paragraph())
    doc.add_page_break()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--md', required=True, help='markdown draft to render')
    ap.add_argument('--out', required=True, help='output .docx path')
    ap.add_argument('--client', required=True, help='client display name for cover and footer')
    ap.add_argument('--doc-title', default='NetSuite Rapid Assessment',
                    help='document title under the client name (default: NetSuite Rapid Assessment)')
    ap.add_argument('--client-logo', default=None, help='client logo image (optional)')
    ap.add_argument('--cover-line', action='append', default=[],
                    help='extra cover line (dates, engagement type); repeatable')
    ap.add_argument('--start-at', default='## Customer Summary',
                    help='first markdown line of the body; everything before it is skipped '
                         '(default: "## Customer Summary", falls back to the first ## heading)')
    ap.add_argument('--no-toc', action='store_true',
                    help='skip the table of contents page (short documents such as executive summaries)')
    args = ap.parse_args()

    src = Path(args.md)
    lines = src.read_text(encoding='utf-8').split('\n')

    # find the body start
    start = next((i for i, ln in enumerate(lines) if ln.strip().startswith(args.start_at)), None)
    if start is None:
        start = next((i for i, ln in enumerate(lines) if ln.startswith('## ')), 0)
        print(f"note: start marker {args.start_at!r} not found; starting at line {start + 1}")
    lines = lines[start:]

    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = BODY
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = MATTE_BLACK

    add_cover(doc, args)
    add_footer(doc, args)
    if not args.no_toc:
        add_toc_page(doc)
    set_updatefields(doc)

    i, n = 0, len(lines)
    while i < n:
        stripped = lines[i].strip()
        if stripped.startswith('```'):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            add_code_block(doc, code_lines)
            continue
        m = re.match(r'^(#{2,4})\s+(.*)$', stripped)
        if m:
            text = m.group(2).strip()
            add_heading(doc, text, len(m.group(1)) - 1, bookmark=slugify(text))
            i += 1
            continue
        if stripped.startswith('|'):
            block = []
            while i < n and lines[i].strip().startswith('|'):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_markdown_table(block))
            continue
        if stripped.startswith(('- ', '* ')):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            add_rich_text(p, stripped[2:].strip())
            i += 1
            continue
        mo = re.match(r'^\d+\.\s+(.*)$', stripped)
        if mo:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(4)
            add_rich_text(p, mo.group(1))
            i += 1
            continue
        if stripped.startswith('*[') and stripped.endswith(']*'):
            p = add_para(doc, '', space_after=8)
            add_rich_text(p, stripped.strip('*'), base_color=DEEP_BLUE)
            for r in p.runs:
                r.italic = True
            i += 1
            continue
        if stripped in ('', '---') or stripped.startswith('# '):
            i += 1
            continue
        p = add_para(doc, '', space_after=8)
        add_rich_text(p, stripped)
        i += 1

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"Saved: {out}")
    print("TOC is a field: run export_pdf.py (updates fields) and check every link in the PDF.")
    return 0


if __name__ == '__main__':
    sys.exit(main())

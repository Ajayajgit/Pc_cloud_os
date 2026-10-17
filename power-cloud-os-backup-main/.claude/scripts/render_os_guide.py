# Render ONBOARDING.md -> branded "How to Work in the Power Cloud OS.docx" + PDF.
# Brand: Montserrat headings, Barlow body, Sky Blue #31A4F3 / Deep Blue #002F59 accents.
import re, os, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = r"G:\Shared drives\Power Cloud OS"
SRC = os.path.join(ROOT, "ONBOARDING.md")
OUT_DOCX = os.path.join(ROOT, "How to Work in the Power Cloud OS.docx")
OUT_PDF = os.path.join(ROOT, "How to Work in the Power Cloud OS.pdf")
LOGO = os.path.join(ROOT, "logos", "Powercloud_Logo-01.png")

SKY = RGBColor(0x31, 0xA4, 0xF3)
DEEP = RGBColor(0x00, 0x2F, 0x59)
INK = RGBColor(0x1E, 0x1E, 0x1E)
MUTED = RGBColor(0x55, 0x67, 0x7B)
GREEN = RGBColor(0x1D, 0x6B, 0x28)

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.9); s.right_margin = Inches(0.9)

st = doc.styles["Normal"]
st.font.name = "Barlow"; st.font.size = Pt(10.5); st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15

def run_inline(p, text, bold=False, italic=False, color=None, mono=False, size=None):
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    if mono: r.font.name = "Consolas"
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    return r

MD_TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`[^`]+`)")
def add_rich(p, text, base_color=None, base_italic=False):
    # strip markdown links [text](url) -> text
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for part in MD_TOKEN.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            run_inline(p, part[2:-2], bold=True, color=base_color, italic=base_italic)
        elif part.startswith("`") and part.endswith("`"):
            run_inline(p, part[1:-1], mono=True, color=DEEP, size=9.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run_inline(p, part[1:-1], italic=True, color=base_color)
        else:
            run_inline(p, part, color=base_color, italic=base_italic)

def heading(text, level):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Montserrat"; r.bold = True
    if level == 1: r.font.size = Pt(22); r.font.color.rgb = DEEP; p.paragraph_format.space_before = Pt(4)
    elif level == 2: r.font.size = Pt(15); r.font.color.rgb = DEEP; p.paragraph_format.space_before = Pt(16)
    else: r.font.size = Pt(11.5); r.font.color.rgb = SKY; p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

# cover block
if os.path.isfile(LOGO):
    lp = doc.add_paragraph(); lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    try: lp.add_run().add_picture(LOGO, width=Inches(1.9))
    except Exception: pass

i = 0
in_code = False
table_buf = []

def flush_table():
    global table_buf
    if not table_buf: return
    rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in table_buf
            if not re.match(r"^\s*\|[\s\-|]+\|\s*$", r)]
    if rows:
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row):
                cp = t.rows[ri].cells[ci].paragraphs[0]
                cp.paragraph_format.space_after = Pt(2)
                if ri == 0:
                    r = cp.add_run(re.sub(r"\*\*", "", cell)); r.bold = True
                    r.font.name = "Montserrat"; r.font.size = Pt(9); r.font.color.rgb = DEEP
                else:
                    add_rich(cp, cell)
                    for r in cp.runs: r.font.size = Pt(9.5)
    table_buf = []

skip_banner = True
while i < len(lines):
    ln = lines[i]
    if ln.startswith("```"):
        flush_table(); in_code = not in_code; i += 1; continue
    if in_code:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.space_after = Pt(0)
        run_inline(p, ln if ln else " ", mono=True, size=9)
        i += 1; continue
    if ln.strip().startswith("|"):
        table_buf.append(ln); i += 1; continue
    flush_table()
    s = ln.strip()
    if not s or s == "---":
        i += 1; continue
    if s.startswith("> **\"") or (s.startswith("> \"") and s.endswith("\"")):
        # sample prompt
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.space_after = Pt(2)
        add_rich(p, s[1:].strip(), base_color=DEEP, base_italic=True)
        i += 1; continue
    if s.startswith(">"):
        if skip_banner and "Retitled and expanded" in s:
            i += 1; continue  # internal draft banner stays out of the distributed copy
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
        add_rich(p, s[1:].strip(), base_color=MUTED, base_italic=True)
        i += 1; continue
    m = re.match(r"^(#{1,3})\s+(.*)$", s)
    if m:
        lvl = len(m.group(1)); text = re.sub(r"\*\*", "", m.group(2))
        if lvl == 1:
            heading(text, 1)
            sub = doc.add_paragraph()
            run_inline(sub, "The Power Cloud workspace guide  ·  Version 2.0 draft  ·  August 12, 2026", color=MUTED, size=10)
            sub.paragraph_format.space_after = Pt(14)
        else:
            heading(text, lvl)
        i += 1; continue
    if s.startswith("- [ ]"):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25); p.paragraph_format.space_after = Pt(2)
        run_inline(p, "\u2610  ", color=SKY)
        add_rich(p, s[5:].strip())
        i += 1; continue
    m = re.match(r"^(\s*)-\s+(.*)$", ln)
    if m:
        p = doc.add_paragraph(style="List Bullet")
        add_rich(p, m.group(2)); p.paragraph_format.space_after = Pt(3)
        i += 1; continue
    m = re.match(r"^(\d+)\.\s+(.*)$", s)
    if m:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25); p.paragraph_format.space_after = Pt(3)
        run_inline(p, m.group(1) + ".  ", bold=True, color=SKY)
        add_rich(p, m.group(2))
        i += 1; continue
    p = doc.add_paragraph(); add_rich(p, s)
    i += 1

flush_table()
doc.save(OUT_DOCX)
print("docx saved:", OUT_DOCX)

# PDF via Word COM
import comtypes.client
word = comtypes.client.CreateObject("Word.Application"); word.Visible = False
try:
    d = word.Documents.Open(OUT_DOCX)
    d.SaveAs(OUT_PDF, FileFormat=17)
    pages = d.ComputeStatistics(2)  # wdStatisticPages
    d.Close(False)
    print(f"pdf saved: {OUT_PDF} ({pages} pages)")
finally:
    word.Quit()
